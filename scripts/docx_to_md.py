"""Make plain-text Markdown mirrors of every .docx in the repo.

Claude (and GitHub search) can't read .docx directly, so each Word file gets a
text copy under catalog/text/<same path>.md. Re-run after adding new reports:

    python3 scripts/docx_to_md.py            # only new/changed files
    python3 scripts/docx_to_md.py --all      # rebuild everything
"""
import sys
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "catalog" / "text"
SKIP_TOP = {"catalog", ".git", ".claude"}


def convert(src: Path) -> str:
    d = docx.Document(src)
    lines = [f"<!-- text mirror of {src.relative_to(ROOT)} — edit the .docx, not this file -->", ""]
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower()
        if style.startswith("heading") or style == "title":
            level = 1 if style == "title" else min(int(style[-1]) if style[-1].isdigit() else 2, 4) + 1
            lines.append("#" * level + " " + t)
        else:
            lines.append(t)
        lines.append("")
    for ti, table in enumerate(d.tables, 1):
        rows = [[c.text.strip().replace("\n", " ").replace("|", "/") for c in r.cells] for r in table.rows]
        if not rows:
            continue
        lines += [f"**Table {ti}**", "", "| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
        lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        lines.append("")
    return "\n".join(lines)


def main():
    rebuild = "--all" in sys.argv
    n = 0
    for src in sorted(ROOT.rglob("*.docx")):
        rel = src.relative_to(ROOT)
        if rel.parts[0] in SKIP_TOP or src.name.startswith("~$"):
            continue
        dst = OUT / rel.with_suffix(".md")
        if not rebuild and dst.exists() and dst.stat().st_mtime >= src.stat().st_mtime:
            continue
        dst.parent.mkdir(parents=True, exist_ok=True)
        try:
            dst.write_text(convert(src))
            n += 1
        except Exception as e:  # corrupt or non-standard docx
            print(f"skip {rel}: {e}")
    print(f"wrote {n} text mirrors to {OUT.relative_to(ROOT)}/")


if __name__ == "__main__":
    main()
